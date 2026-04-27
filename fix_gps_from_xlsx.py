"""
Script: fix_gps_from_xlsx.py
Update latitud + longitud dalam MaklumatAsas menggunakan true_fasiliti.xlsx.

Nota: Dalam fail xlsx, label TERBALIK (sama seperti PDF asal):
  Kolum "Longitude" → nilai ~1-7    → simpan sebagai latitud
  Kolum "Latitude"  → nilai ~100-119 → simpan sebagai longitud

Jalankan: python fix_gps_from_xlsx.py
"""

import os
import sys
import django
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ybtmk_project.settings')
django.setup()

from home.models import MaklumatAsas

XLSX_PATH = r'c:\Users\rahma\Downloads\true_fasiliti.xlsx'

def main():
    print(f'Membuka: {XLSX_PATH}')
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    # Baca headers → cari index kolum yang diperlukan
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    print(f'Headers: {headers}')

    try:
        idx_nama = headers.index('Nama_Fasiliti') + 1
        idx_lng  = headers.index('Longitude') + 1   # nilai ~1-7  → latitud
        idx_lat  = headers.index('Latitude') + 1    # nilai ~100-119 → longitud
    except ValueError as e:
        print(f'ERROR: Kolum tidak dijumpai — {e}')
        sys.exit(1)

    print(f'Kolum: Nama_Fasiliti={idx_nama}, Longitude(→latitud)={idx_lng}, Latitude(→longitud)={idx_lat}')
    print()

    # Bina lookup DB: nama_lower → MaklumatAsas object
    print('Memuat DB records...')
    db_lookup = {}
    for obj in MaklumatAsas.objects.all():
        key = obj.nama_fasiliti.strip().lower()
        db_lookup[key] = obj
    print(f'  {len(db_lookup)} records dalam DB')
    print()

    updated   = 0
    skipped   = 0
    notfound  = 0
    bad_gps   = 0
    bulk_objs = []

    total_rows = ws.max_row - 1  # exclude header
    for r in range(2, ws.max_row + 1):
        nama    = ws.cell(r, idx_nama).value
        raw_lat = ws.cell(r, idx_lng).value   # Longitude col = latitud sebenar
        raw_lng = ws.cell(r, idx_lat).value   # Latitude col  = longitud sebenar

        if not nama:
            skipped += 1
            continue

        # Convert ke float
        try:
            lat = float(raw_lat) if raw_lat is not None else None
            lng = float(raw_lng) if raw_lng is not None else None
        except (ValueError, TypeError):
            lat = lng = None

        # Sanity check — Malaysia range
        if lat is not None and lng is not None:
            if not (0.8 <= lat <= 7.5 and 99.5 <= lng <= 119.5):
                print(f'  [BAD GPS] "{nama}" lat={lat} lng={lng} — skip')
                bad_gps += 1
                continue

        # Padanan nama
        key = nama.strip().lower()
        obj = db_lookup.get(key)

        # Cuba strip negeri suffix jika tiada padanan terus
        if obj is None:
            key2 = key.split(',')[0].strip()
            obj = db_lookup.get(key2)

        if obj is None:
            notfound += 1
            continue

        if obj.latitud == lat and obj.longitud == lng:
            skipped += 1
            continue

        obj.latitud  = lat
        obj.longitud = lng
        bulk_objs.append(obj)

        # Batch save setiap 500
        if len(bulk_objs) >= 500:
            MaklumatAsas.objects.bulk_update(bulk_objs, ['latitud', 'longitud'])
            updated += len(bulk_objs)
            print(f'  ... {updated}/{total_rows} dikemaskini')
            bulk_objs = []

    # Baki
    if bulk_objs:
        MaklumatAsas.objects.bulk_update(bulk_objs, ['latitud', 'longitud'])
        updated += len(bulk_objs)

    print()
    print(f'SELESAI:')
    print(f'  Dikemaskini : {updated}')
    print(f'  Sama (skip) : {skipped}')
    print(f'  Tiada dalam DB : {notfound}')
    print(f'  GPS tidak valid: {bad_gps}')

if __name__ == '__main__':
    main()
